from pathlib import Path
import json
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook


ARQUIVO_EXCEL = Path("clientes.xlsx")
ARQUIVO_LOGO_PNG = Path("logo.png")
ARQUIVO_LOGO_PPM = Path("logo.ppm")
ARQUIVO_CONFIG = Path("config.json")
COLUNAS = [
    "ID",
    "Nome",
    "Email",
    "Telefone",
    "CPF",
    "Cidade",
    "Time atual",
    "Pe dominante",
    "Numero da camisa",
    "Posicao",
    "Gols",
    "Penaltis defendidos",
    "Foto",
]
POSICOES = ["Goleiro", "Zagueiro", "Lateral", "Volante", "Meia", "Atacante"]
PES_DOMINANTES = ["Direito", "Esquerdo", "Ambidestro"]
TEMAS = {
    "claro": {
        "root_bg": "#f3f6fb",
        "header_bg": "#eef4fb",
        "body_bg": "#f3f6fb",
        "card_bg": "#ffffff",
        "title_fg": "#12324a",
        "subtitle_fg": "#567086",
        "search_fg": "#4c6172",
        "section_fg": "#12324a",
        "label_fg": "#12324a",
        "entry_bg": "#f7fafc",
        "entry_border": "#d7e3ef",
        "tree_bg": "#fdfefe",
        "tree_field_bg": "#ffffff",
        "tree_heading_bg": "#dbe8f4",
        "tree_heading_fg": "#12324a",
        "tree_selected_bg": "#9ec5fe",
        "tree_selected_fg": "#0f2233",
        "row_even": "#ffffff",
        "row_odd": "#f4f8fc",
    },
    "escuro": {
        "root_bg": "#111827",
        "header_bg": "#172033",
        "body_bg": "#111827",
        "card_bg": "#1b2538",
        "title_fg": "#f3f7ff",
        "subtitle_fg": "#a9b7cb",
        "search_fg": "#c3d0e0",
        "section_fg": "#e7eef9",
        "label_fg": "#e7eef9",
        "entry_bg": "#243146",
        "entry_border": "#32435d",
        "tree_bg": "#1a2435",
        "tree_field_bg": "#1a2435",
        "tree_heading_bg": "#26354c",
        "tree_heading_fg": "#edf3ff",
        "tree_selected_bg": "#315a9a",
        "tree_selected_fg": "#ffffff",
        "row_even": "#1a2435",
        "row_odd": "#202c40",
    },
}


def formatar_cpf(cpf):
    digitos = "".join(caractere for caractere in cpf if caractere.isdigit())[:11]

    if len(digitos) <= 3:
        return digitos
    if len(digitos) <= 6:
        return f"{digitos[:3]}.{digitos[3:]}"
    if len(digitos) <= 9:
        return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:]}"
    return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"


def formatar_telefone(telefone):
    digitos = "".join(caractere for caractere in telefone if caractere.isdigit())[:11]

    if not digitos:
        return ""
    if len(digitos) <= 2:
        return f"({digitos}"
    if len(digitos) <= 6:
        return f"({digitos[:2]}) {digitos[2:]}"
    if len(digitos) <= 10:
        return f"({digitos[:2]}) {digitos[2:6]}-{digitos[6:]}"
    return f"({digitos[:2]}) {digitos[2:7]}-{digitos[7:]}"


def criar_nova_planilha():
    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Clientes"
    planilha.append(COLUNAS)
    workbook.save(ARQUIVO_EXCEL)


def inicializar_planilha():
    if ARQUIVO_EXCEL.exists():
        garantir_colunas_planilha()
        migrar_dados_legado()
        vincular_foto_padrao_yuri()
        return

    criar_nova_planilha()
    migrar_dados_legado()
    vincular_foto_padrao_yuri()


def abrir_planilha():
    try:
        workbook = load_workbook(ARQUIVO_EXCEL)
    except (BadZipFile, OSError):
        arquivo_backup = ARQUIVO_EXCEL.with_name(f"{ARQUIVO_EXCEL.stem}_corrompido.xlsx")

        if arquivo_backup.exists():
            arquivo_backup.unlink()

        ARQUIVO_EXCEL.replace(arquivo_backup)
        criar_nova_planilha()
        workbook = load_workbook(ARQUIVO_EXCEL)

    return workbook, workbook["Clientes"]


def garantir_colunas_planilha():
    workbook, planilha = abrir_planilha()
    alterou_planilha = False

    for indice, coluna in enumerate(COLUNAS, start=1):
        if planilha.cell(row=1, column=indice).value != coluna:
            planilha.cell(row=1, column=indice).value = coluna
            if indice > planilha.max_column:
                planilha.cell(row=2, column=indice).value = ""
                planilha.cell(row=2, column=indice).value = None
            alterou_planilha = True

    if planilha.max_row == 1 and planilha.max_column < len(COLUNAS):
        for indice in range(1, len(COLUNAS) + 1):
            planilha.cell(row=1, column=indice).value = COLUNAS[indice - 1]
        alterou_planilha = True

    if alterou_planilha:
        workbook.save(ARQUIVO_EXCEL)


def migrar_dados_legado():
    workbook, planilha = abrir_planilha()
    alterou_planilha = False

    for linha in range(2, planilha.max_row + 1):
        valor_posicao = planilha.cell(row=linha, column=10).value
        valor_gols = planilha.cell(row=linha, column=11).value
        valor_penaltis = planilha.cell(row=linha, column=12).value
        valor_foto = planilha.cell(row=linha, column=13).value

        for coluna_origem, valor in ((10, valor_posicao), (11, valor_gols), (12, valor_penaltis)):
            if isinstance(valor, str) and valor.lower().endswith((".png", ".gif")):
                if not valor_foto:
                    planilha.cell(row=linha, column=13).value = valor
                    valor_foto = valor
                planilha.cell(row=linha, column=coluna_origem).value = ""
                alterou_planilha = True

        valor_posicao_atual = planilha.cell(row=linha, column=10).value
        valor_gols_atual = planilha.cell(row=linha, column=11).value
        valor_penaltis_atual = planilha.cell(row=linha, column=12).value

        if valor_posicao_atual == "Goleiro" and valor_gols_atual and not valor_penaltis_atual:
            planilha.cell(row=linha, column=12).value = valor_gols_atual
            planilha.cell(row=linha, column=11).value = None
            alterou_planilha = True

    if alterou_planilha:
        workbook.save(ARQUIVO_EXCEL)


def proximo_id(planilha):
    ids = []
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if linha[0] is not None:
            ids.append(int(linha[0]))
    return max(ids, default=0) + 1


def listar_clientes():
    _, planilha = abrir_planilha()
    clientes = []

    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if linha[0] is None:
            continue
        clientes.append(
            (
                linha[0],
                linha[1] or "",
                linha[2] or "",
                linha[3] or "",
                linha[4] or "",
                linha[5] or "",
                linha[6] or "",
                linha[7] or "",
                linha[8] or "",
                linha[9] or "",
                linha[10] or "",
                linha[11] or "",
                linha[12] or "",
            )
        )

    return clientes


def buscar_cliente_por_id(cliente_id):
    for cliente in listar_clientes():
        if int(cliente[0]) == int(cliente_id):
            return cliente
    return None


def buscar_linha_por_id(planilha, cliente_id):
    for linha in range(2, planilha.max_row + 1):
        if planilha.cell(row=linha, column=1).value == cliente_id:
            return linha
    return None


def cadastrar_cliente(
    nome,
    email,
    telefone,
    cpf,
    cidade,
    time_atual,
    pe_dominante,
    numero_camisa,
    posicao,
    desempenho,
    foto="",
):
    workbook, planilha = abrir_planilha()
    cliente_id = proximo_id(planilha)
    gols = desempenho if posicao != "Goleiro" else ""
    penaltis_defendidos = desempenho if posicao == "Goleiro" else ""
    planilha.append(
        [
            cliente_id,
            nome,
            email,
            telefone,
            cpf,
            cidade,
            time_atual,
            pe_dominante,
            numero_camisa,
            posicao,
            gols,
            penaltis_defendidos,
            foto,
        ]
    )
    workbook.save(ARQUIVO_EXCEL)
    return cliente_id


def atualizar_cliente(
    cliente_id,
    nome,
    email,
    telefone,
    cpf,
    cidade,
    time_atual,
    pe_dominante,
    numero_camisa,
    posicao,
    desempenho,
    foto="",
):
    workbook, planilha = abrir_planilha()
    linha = buscar_linha_por_id(planilha, cliente_id)

    if linha is None:
        return False

    planilha.cell(row=linha, column=2).value = nome
    planilha.cell(row=linha, column=3).value = email
    planilha.cell(row=linha, column=4).value = telefone
    planilha.cell(row=linha, column=5).value = cpf
    planilha.cell(row=linha, column=6).value = cidade
    planilha.cell(row=linha, column=7).value = time_atual
    planilha.cell(row=linha, column=8).value = pe_dominante
    planilha.cell(row=linha, column=9).value = numero_camisa
    planilha.cell(row=linha, column=10).value = posicao
    planilha.cell(row=linha, column=11).value = desempenho if posicao != "Goleiro" else ""
    planilha.cell(row=linha, column=12).value = desempenho if posicao == "Goleiro" else ""
    planilha.cell(row=linha, column=13).value = foto
    workbook.save(ARQUIVO_EXCEL)
    return True


def atualizar_foto_por_nome(nome_cliente, foto):
    workbook, planilha = abrir_planilha()

    for linha in range(2, planilha.max_row + 1):
        nome_planilha = str(planilha.cell(row=linha, column=2).value or "").strip().lower()
        if nome_planilha == nome_cliente.strip().lower():
            planilha.cell(row=linha, column=13).value = foto
            workbook.save(ARQUIVO_EXCEL)
            return True

    return False


def excluir_cliente(cliente_id):
    workbook, planilha = abrir_planilha()
    linha = buscar_linha_por_id(planilha, cliente_id)

    if linha is None:
        return False

    planilha.delete_rows(linha)
    workbook.save(ARQUIVO_EXCEL)
    return True


def vincular_foto_padrao_yuri():
    arquivo_foto = Path("yuri_alberto.png")
    if arquivo_foto.exists():
        atualizar_foto_por_nome("Yuri alberto", arquivo_foto.name)


def carregar_configuracoes():
    if not ARQUIVO_CONFIG.exists():
        return {"tema": "claro"}

    try:
        with ARQUIVO_CONFIG.open("r", encoding="utf-8") as arquivo:
            configuracoes = json.load(arquivo)
    except (json.JSONDecodeError, OSError):
        return {"tema": "claro"}

    tema = configuracoes.get("tema", "claro")
    if tema not in TEMAS:
        tema = "claro"

    return {"tema": tema}


def salvar_configuracoes(configuracoes):
    with ARQUIVO_CONFIG.open("w", encoding="utf-8") as arquivo:
        json.dump(configuracoes, arquivo, ensure_ascii=False, indent=2)


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("CRUD de Jogadores com Excel")
        self.root.geometry("1280x900")
        self.root.minsize(1100, 760)
        self.root.resizable(True, True)
        self.root.configure(bg="#f3f6fb")
        self.configurar_janela_inicial()

        self.id_cliente = None
        self.logo_imagem = None
        self.configuracoes = carregar_configuracoes()

        self.nome_var = tk.StringVar()
        self.email_var = tk.StringVar()
        self.telefone_var = tk.StringVar()
        self.cpf_var = tk.StringVar()
        self.cidade_var = tk.StringVar()
        self.time_atual_var = tk.StringVar()
        self.pe_dominante_var = tk.StringVar()
        self.numero_camisa_var = tk.StringVar()
        self.posicao_var = tk.StringVar()
        self.desempenho_var = tk.StringVar()
        self.label_desempenho_var = tk.StringVar(value="Gols")
        self.foto_var = tk.StringVar()
        self.busca_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Sistema pronto.")
        self.atualizando_telefone = False
        self.atualizando_cpf = False
        self.foto_cliente_imagem = None
        self.ordenar_por = "ID"
        self.ordem_reversa = False
        self.tema_atual = self.configuracoes["tema"]
        self.tema_var = tk.StringVar(value="Modo escuro")

        self.configurar_estilos()
        self.criar_widgets()
        self.telefone_var.trace_add("write", self.aplicar_mascara_telefone)
        self.cpf_var.trace_add("write", self.aplicar_mascara_cpf)
        self.posicao_var.trace_add("write", self.atualizar_label_desempenho)
        self.carregar_clientes()
        self.tema_var.set("Modo claro" if self.tema_atual == "escuro" else "Modo escuro")

    def configurar_janela_inicial(self):
        try:
            self.root.state("zoomed")
        except tk.TclError:
            pass

    def configurar_estilos(self):
        style = ttk.Style()
        style.theme_use("clam")
        self.aplicar_tema()

    def aplicar_tema(self):
        style = ttk.Style()
        cores = TEMAS[self.tema_atual]

        self.root.configure(bg=cores["root_bg"])
        style.configure(".", font=("Segoe UI", 10))
        style.configure("Header.TFrame", background=cores["header_bg"])
        style.configure("Body.TFrame", background=cores["body_bg"])
        style.configure("Card.TFrame", background=cores["card_bg"], relief="flat")
        style.configure(
            "Title.TLabel",
            background=cores["header_bg"],
            foreground=cores["title_fg"],
            font=("Georgia", 22, "bold"),
        )
        style.configure(
            "Subtitle.TLabel",
            background=cores["header_bg"],
            foreground=cores["subtitle_fg"],
            font=("Segoe UI", 10),
        )
        style.configure(
            "SearchLabel.TLabel",
            background=cores["card_bg"],
            foreground=cores["search_fg"],
            font=("Segoe UI", 10),
        )
        style.configure(
            "Section.TLabelframe",
            background=cores["card_bg"],
            borderwidth=0,
            relief="solid",
        )
        style.configure(
            "Section.TLabelframe.Label",
            background=cores["card_bg"],
            foreground=cores["section_fg"],
            font=("Segoe UI", 11, "bold"),
        )
        style.configure(
            "TLabel",
            background=cores["card_bg"],
            foreground=cores["label_fg"],
        )
        style.configure(
            "TEntry",
            fieldbackground=cores["entry_bg"],
            bordercolor=cores["entry_border"],
            lightcolor=cores["entry_border"],
            darkcolor=cores["entry_border"],
            padding=8,
            foreground=cores["label_fg"],
        )
        style.configure(
            "Treeview",
            background=cores["tree_bg"],
            fieldbackground=cores["tree_field_bg"],
            rowheight=34,
            font=("Segoe UI", 10),
            borderwidth=0,
            foreground=cores["label_fg"],
        )
        style.configure(
            "Treeview.Heading",
            background=cores["tree_heading_bg"],
            foreground=cores["tree_heading_fg"],
            font=("Segoe UI", 10, "bold"),
            relief="flat",
        )
        style.map(
            "Treeview",
            background=[("selected", cores["tree_selected_bg"])],
            foreground=[("selected", cores["tree_selected_fg"])],
        )

        style.configure("Primary.TButton", background="#167c80", foreground="#ffffff", padding=(14, 9))
        style.map("Primary.TButton", background=[("active", "#116468")])
        style.configure("Info.TButton", background="#2f6fed", foreground="#ffffff", padding=(14, 9))
        style.map("Info.TButton", background=[("active", "#2458bb")])
        style.configure("Danger.TButton", background="#c4543d", foreground="#ffffff", padding=(14, 9))
        style.map("Danger.TButton", background=[("active", "#9d4331")])
        style.configure("Neutral.TButton", background="#738496", foreground="#ffffff", padding=(14, 9))
        style.map("Neutral.TButton", background=[("active", "#5c6978")])

        if hasattr(self, "tabela"):
            self.tabela.tag_configure("linha_par", background=cores["row_even"], foreground=cores["label_fg"])
            self.tabela.tag_configure("linha_impar", background=cores["row_odd"], foreground=cores["label_fg"])

    def criar_widgets(self):
        self.criar_cabecalho()

        frame_conteudo = ttk.Frame(self.root, style="Body.TFrame")
        frame_conteudo.pack(fill="both", expand=True)
        frame_conteudo.columnconfigure(0, weight=1)
        frame_conteudo.rowconfigure(2, weight=1)

        frame_busca = ttk.LabelFrame(frame_conteudo, text="Busca", padding=18, style="Section.TLabelframe")
        frame_busca.pack(fill="x", padx=16, pady=(12, 8))

        ttk.Label(
            frame_busca,
            text="Pesquisar por nome, CPF ou cidade",
            style="SearchLabel.TLabel",
        ).pack(side="left", padx=(0, 12))
        ttk.Entry(frame_busca, textvariable=self.busca_var, width=42).pack(side="left")
        ttk.Button(frame_busca, text="Buscar", command=self.aplicar_busca, style="Info.TButton").pack(
            side="left", padx=(12, 8)
        )
        ttk.Button(frame_busca, text="Limpar busca", command=self.limpar_busca, style="Neutral.TButton").pack(
            side="left"
        )

        frame_formulario = ttk.LabelFrame(
            frame_conteudo, text="Dados do jogador", padding=22, style="Section.TLabelframe"
        )
        frame_formulario.pack(fill="x", padx=16, pady=(16, 8))
        frame_formulario.columnconfigure(0, weight=1)
        frame_formulario.columnconfigure(1, weight=1)
        frame_formulario.columnconfigure(2, weight=1)
        frame_formulario.columnconfigure(3, weight=1)
        frame_formulario.columnconfigure(4, weight=1)
        frame_formulario.columnconfigure(5, weight=1)

        ttk.Label(frame_formulario, text="Nome").grid(row=0, column=0, sticky="w")
        self.nome_entry = ttk.Entry(frame_formulario, textvariable=self.nome_var, width=20)
        self.nome_entry.grid(
            row=1, column=0, padx=(0, 14), pady=(0, 16), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Telefone").grid(row=0, column=1, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.telefone_var, width=20).grid(
            row=1, column=1, padx=(0, 14), pady=(0, 16), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Email").grid(row=0, column=2, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.email_var, width=20).grid(
            row=1, column=2, pady=(0, 16), sticky="ew"
        )

        ttk.Label(frame_formulario, text="CPF").grid(row=2, column=0, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.cpf_var, width=20).grid(
            row=3, column=0, padx=(0, 14), pady=(0, 18), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Cidade").grid(row=2, column=1, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.cidade_var, width=20).grid(
            row=3, column=1, padx=(0, 14), pady=(0, 18), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Time atual").grid(row=2, column=2, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.time_atual_var, width=20).grid(
            row=3, column=2, padx=(0, 14), pady=(0, 18), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Pé dominante").grid(row=2, column=3, sticky="w")
        ttk.Combobox(
            frame_formulario,
            textvariable=self.pe_dominante_var,
            values=PES_DOMINANTES,
            state="readonly",
            width=18,
        ).grid(row=3, column=3, padx=(0, 14), pady=(0, 18), sticky="ew")

        ttk.Label(frame_formulario, text="Número da camisa").grid(row=2, column=4, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.numero_camisa_var, width=20).grid(
            row=3, column=4, padx=(0, 14), pady=(0, 18), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Posição").grid(row=4, column=0, sticky="w")
        ttk.Combobox(
            frame_formulario,
            textvariable=self.posicao_var,
            values=POSICOES,
            state="readonly",
            width=18,
        ).grid(row=5, column=0, padx=(0, 14), pady=(0, 18), sticky="ew")

        ttk.Label(frame_formulario, textvariable=self.label_desempenho_var).grid(row=4, column=1, sticky="w")
        ttk.Entry(frame_formulario, textvariable=self.desempenho_var, width=20).grid(
            row=5, column=1, padx=(0, 14), pady=(0, 18), sticky="ew"
        )

        ttk.Label(frame_formulario, text="Foto").grid(row=4, column=2, sticky="w")
        ttk.Button(
            frame_formulario,
            text="Selecionar foto",
            command=self.selecionar_foto,
            style="Info.TButton",
        ).grid(row=5, column=2, sticky="w")

        self.label_foto = ttk.Label(frame_formulario, text="Sem foto", anchor="center")
        self.label_foto.grid(row=0, column=6, rowspan=6, padx=(10, 0), sticky="n")

        frame_botoes = ttk.Frame(frame_formulario, style="Card.TFrame")
        frame_botoes.grid(row=6, column=0, columnspan=7, sticky="e")

        ttk.Button(frame_botoes, text="Cadastrar", command=self.cadastrar, style="Primary.TButton").pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(frame_botoes, text="Atualizar", command=self.atualizar, style="Info.TButton").pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(frame_botoes, text="Excluir", command=self.excluir, style="Danger.TButton").pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(frame_botoes, text="Limpar", command=self.limpar_campos, style="Neutral.TButton").pack(
            side="left"
        )

        frame_tabela = ttk.LabelFrame(
            frame_conteudo, text="Jogadores cadastrados", padding=18, style="Section.TLabelframe"
        )
        frame_tabela.pack(fill="both", expand=True, padx=16, pady=(0, 16))
        frame_tabela.columnconfigure(0, weight=1)
        frame_tabela.rowconfigure(0, weight=1)

        colunas = (
            "ID",
            "Nome",
            "Email",
            "Telefone",
            "CPF",
            "Cidade",
            "Time atual",
            "Pé dominante",
            "Número da camisa",
            "Posição",
            "Gols",
            "Penaltis defendidos",
        )
        self.tabela = ttk.Treeview(frame_tabela, columns=colunas, show="headings", height=26)

        for coluna in colunas:
            self.tabela.heading(coluna, text=coluna, command=lambda c=coluna: self.ordenar_tabela(c))

        self.tabela.column("ID", width=80, anchor="center", stretch=False)
        self.tabela.column("Nome", width=320, stretch=True)
        self.tabela.column("Email", width=380, stretch=True)
        self.tabela.column("Telefone", width=150, stretch=False)
        self.tabela.column("CPF", width=180, stretch=False)
        self.tabela.column("Cidade", width=180, stretch=True)
        self.tabela.column("Time atual", width=180, stretch=True)
        self.tabela.column("Pé dominante", width=130, stretch=False)
        self.tabela.column("Número da camisa", width=140, stretch=False)
        self.tabela.column("Posição", width=150, stretch=False)
        self.tabela.column("Gols", width=160, stretch=False)
        self.tabela.column("Penaltis defendidos", width=160, stretch=False)

        self.tabela.tag_configure("linha_par", background="#eee7e7")
        self.tabela.tag_configure("linha_impar", background="#f4f8fc")

        barra_rolagem = ttk.Scrollbar(
            frame_tabela, orient="vertical", command=self.tabela.yview
        )
        barra_rolagem_horizontal = ttk.Scrollbar(
            frame_tabela, orient="horizontal", command=self.tabela.xview
        )
        self.tabela.configure(
            yscrollcommand=barra_rolagem.set,
            xscrollcommand=barra_rolagem_horizontal.set,
        )

        self.tabela.grid(row=0, column=0, sticky="nsew")
        barra_rolagem.grid(row=0, column=1, sticky="ns")
        barra_rolagem_horizontal.grid(row=1, column=0, sticky="ew")

        self.tabela.bind("<<TreeviewSelect>>", self.selecionar_cliente)
        self.tabela.bind("<Double-1>", self.editar_por_duplo_clique)
        self.tabela.bind("<MouseWheel>", self.rolar_tabela_mouse)

        frame_status = ttk.Frame(self.root, style="Header.TFrame", padding=(16, 0, 16, 12))
        frame_status.pack(fill="x")
        ttk.Label(frame_status, textvariable=self.status_var, style="Subtitle.TLabel").pack(anchor="w")

    def criar_cabecalho(self):
        frame_cabecalho = ttk.Frame(self.root, padding=(16, 16, 16, 0), style="Header.TFrame")
        frame_cabecalho.pack(fill="x")

        frame_cabecalho.columnconfigure(0, weight=0)
        frame_cabecalho.columnconfigure(1, weight=1)
        frame_cabecalho.columnconfigure(2, weight=0)

        frame_logo = ttk.Frame(frame_cabecalho, style="Header.TFrame")
        frame_logo.grid(row=0, column=0, sticky="w")

        frame_titulo = ttk.Frame(frame_cabecalho, style="Header.TFrame")
        frame_titulo.grid(row=0, column=1, sticky="ew")

        frame_espaco = ttk.Frame(frame_cabecalho, width=120, style="Header.TFrame")
        frame_espaco.grid(row=0, column=2, sticky="e")

        ttk.Button(
            frame_espaco,
            textvariable=self.tema_var,
            command=self.alternar_tema,
            style="Neutral.TButton",
        ).pack(anchor="e")

        caminho_logo = None
        if ARQUIVO_LOGO_PNG.exists():
            caminho_logo = ARQUIVO_LOGO_PNG
        elif ARQUIVO_LOGO_PPM.exists():
            caminho_logo = ARQUIVO_LOGO_PPM

        if caminho_logo is not None:
            try:
                self.logo_imagem = tk.PhotoImage(file=str(caminho_logo))
                self.root.iconphoto(True, self.logo_imagem)
                ttk.Label(frame_logo, image=self.logo_imagem).pack(side="left")
            except tk.TclError:
                ttk.Label(frame_logo, text="Logo nao pode ser carregado.").pack(side="left")
        else:
            ttk.Label(
                frame_logo,
                text="Adicione um arquivo logo.png ou logo.ppm na pasta do projeto.",
            ).pack(side="left")

        ttk.Label(
            frame_titulo,
            text="Cadastro de Jogadores",
            style="Title.TLabel",
        ).pack()
        ttk.Label(
            frame_titulo,
            text="Gerencie seus jogadores com Excel de forma simples e organizada",
            style="Subtitle.TLabel",
        ).pack(pady=(4, 0))

    def carregar_clientes(self):
        termo_busca = self.busca_var.get().strip().lower()
        clientes = []

        for item in self.tabela.get_children():
            self.tabela.delete(item)

        for cliente in listar_clientes():
            if termo_busca:
                texto_cliente = " ".join(str(valor).lower() for valor in cliente[1:])
                if termo_busca not in texto_cliente:
                    continue
            clientes.append(cliente)

        clientes = self.clientes_ordenados(clientes)

        for cliente in clientes:
            indice = len(self.tabela.get_children())
            tag = "linha_par" if indice % 2 == 0 else "linha_impar"
            self.tabela.insert("", "end", values=cliente[:12], tags=(tag,))

        quantidade = len(self.tabela.get_children())
        if termo_busca:
            self.status_var.set(f"{quantidade} jogador(es) encontrado(s) para '{self.busca_var.get().strip()}'.")
        else:
            self.status_var.set(f"{quantidade} jogador(es) listado(s).")

    def clientes_ordenados(self, clientes):
        indice_coluna = {
            "ID": 0,
            "Nome": 1,
            "Email": 2,
            "Telefone": 3,
            "CPF": 4,
            "Cidade": 5,
            "Time atual": 6,
            "Pé dominante": 7,
            "Número da camisa": 8,
            "Posição": 9,
            "Gols": 10,
            "Penaltis defendidos": 11,
        }[self.ordenar_por]

        if self.ordenar_por == "ID":
            chave = lambda cliente: int(cliente[indice_coluna])
        else:
            chave = lambda cliente: str(cliente[indice_coluna]).lower()

        return sorted(clientes, key=chave, reverse=self.ordem_reversa)

    def ordenar_tabela(self, coluna):
        if self.ordenar_por == coluna:
            self.ordem_reversa = not self.ordem_reversa
        else:
            self.ordenar_por = coluna
            self.ordem_reversa = False

        self.carregar_clientes()
        ordem = "decrescente" if self.ordem_reversa else "crescente"
        self.status_var.set(f"Tabela ordenada por {coluna.lower()} em ordem {ordem}.")

    def aplicar_busca(self):
        self.carregar_clientes()

    def limpar_busca(self):
        self.busca_var.set("")
        self.carregar_clientes()

    def alternar_tema(self):
        self.tema_atual = "escuro" if self.tema_atual == "claro" else "claro"
        self.tema_var.set("Modo claro" if self.tema_atual == "escuro" else "Modo escuro")
        self.configuracoes["tema"] = self.tema_atual
        salvar_configuracoes(self.configuracoes)
        self.aplicar_tema()
        self.carregar_clientes()
        self.atualizar_preview_foto()
        self.status_var.set(
            "Tema escuro ativado." if self.tema_atual == "escuro" else "Tema claro ativado."
        )

    def validar_campos(self):
        if not self.nome_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o nome do jogador.")
            return False

        if not self.email_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o email do jogador.")
            return False

        if not self.telefone_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o telefone do jogador.")
            return False

        telefone = "".join(
            caractere for caractere in self.telefone_var.get() if caractere.isdigit()
        )
        if len(telefone) not in (10, 11):
            messagebox.showwarning("Aviso", "O telefone deve conter 10 ou 11 digitos.")
            return False

        if not self.cpf_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o CPF do jogador.")
            return False

        cpf = "".join(caractere for caractere in self.cpf_var.get() if caractere.isdigit())
        if len(cpf) != 11:
            messagebox.showwarning("Aviso", "O CPF deve conter 11 digitos.")
            return False

        if not self.cidade_var.get().strip():
            messagebox.showwarning("Aviso", "Informe a cidade do jogador.")
            return False

        if not self.time_atual_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o time atual do jogador.")
            return False

        if not self.pe_dominante_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o pé dominante do jogador.")
            return False

        if not self.numero_camisa_var.get().strip():
            messagebox.showwarning("Aviso", "Informe o número da camisa.")
            return False

        if not self.numero_camisa_var.get().strip().isdigit():
            messagebox.showwarning("Aviso", "O número da camisa deve ser numerico.")
            return False

        if not self.posicao_var.get().strip():
            messagebox.showwarning("Aviso", "Informe a posição do jogador.")
            return False

        if not self.desempenho_var.get().strip():
            messagebox.showwarning("Aviso", f"Informe {self.label_desempenho_var.get().lower()}.")
            return False

        if not self.desempenho_var.get().strip().isdigit():
            messagebox.showwarning("Aviso", f"{self.label_desempenho_var.get()} deve ser numerico.")
            return False

        return True

    def atualizar_label_desempenho(self, *_args):
        if self.posicao_var.get() == "Goleiro":
            self.label_desempenho_var.set("Penaltis defendidos")
        else:
            self.label_desempenho_var.set("Gols")

    def aplicar_mascara_cpf(self, *_args):
        if self.atualizando_cpf:
            return

        self.atualizando_cpf = True
        self.cpf_var.set(formatar_cpf(self.cpf_var.get()))
        self.atualizando_cpf = False

    def aplicar_mascara_telefone(self, *_args):
        if self.atualizando_telefone:
            return

        self.atualizando_telefone = True
        self.telefone_var.set(formatar_telefone(self.telefone_var.get()))
        self.atualizando_telefone = False

    def atualizar_preview_foto(self):
        caminho_foto = self.foto_var.get().strip()
        cores = TEMAS[self.tema_atual]
        self.label_foto.configure(background=cores["card_bg"], foreground=cores["label_fg"])

        if not caminho_foto:
            self.foto_cliente_imagem = None
            self.label_foto.configure(image="", text="Sem foto")
            return

        caminho = Path(caminho_foto)
        if not caminho.is_absolute():
            caminho = Path.cwd() / caminho

        if not caminho.exists():
            self.foto_cliente_imagem = None
            self.label_foto.configure(image="", text="Foto nao encontrada")
            return

        try:
            self.foto_cliente_imagem = tk.PhotoImage(file=str(caminho))
            self.label_foto.configure(image=self.foto_cliente_imagem, text="")
        except tk.TclError:
            self.foto_cliente_imagem = None
            self.label_foto.configure(image="", text="Formato invalido")

    def selecionar_foto(self):
        arquivo = filedialog.askopenfilename(
            title="Selecionar foto",
            filetypes=[
                ("Imagens PNG", "*.png"),
                ("Imagens GIF", "*.gif"),
                ("Todos os arquivos", "*.*"),
            ],
        )
        if not arquivo:
            return

        caminho = Path(arquivo)
        try:
            caminho_relativo = caminho.relative_to(Path.cwd())
            self.foto_var.set(str(caminho_relativo))
        except ValueError:
            self.foto_var.set(str(caminho))

        self.atualizar_preview_foto()

    def cadastrar(self):
        if not self.validar_campos():
            return

        cliente_id = cadastrar_cliente(
            self.nome_var.get().strip(),
            self.email_var.get().strip(),
            self.telefone_var.get().strip(),
            self.cpf_var.get().strip(),
            self.cidade_var.get().strip(),
            self.time_atual_var.get().strip(),
            self.pe_dominante_var.get().strip(),
            self.numero_camisa_var.get().strip(),
            self.posicao_var.get().strip(),
            self.desempenho_var.get().strip(),
            self.foto_var.get().strip(),
        )
        self.carregar_clientes()
        self.limpar_campos()
        self.status_var.set(f"Jogador {cliente_id} cadastrado com sucesso.")
        messagebox.showinfo("Sucesso", f"Jogador cadastrado com ID {cliente_id}.")

    def atualizar(self):
        if self.id_cliente is None:
            messagebox.showwarning("Aviso", "Selecione um jogador na tabela.")
            return

        if not self.validar_campos():
            return

        atualizado = atualizar_cliente(
            self.id_cliente,
            self.nome_var.get().strip(),
            self.email_var.get().strip(),
            self.telefone_var.get().strip(),
            self.cpf_var.get().strip(),
            self.cidade_var.get().strip(),
            self.time_atual_var.get().strip(),
            self.pe_dominante_var.get().strip(),
            self.numero_camisa_var.get().strip(),
            self.posicao_var.get().strip(),
            self.desempenho_var.get().strip(),
            self.foto_var.get().strip(),
        )

        if not atualizado:
            messagebox.showerror("Erro", "Jogador nao encontrado.")
            return

        self.carregar_clientes()
        self.limpar_campos()
        self.status_var.set("Jogador atualizado com sucesso.")
        messagebox.showinfo("Sucesso", "Jogador atualizado com sucesso.")

    def excluir(self):
        if self.id_cliente is None:
            messagebox.showwarning("Aviso", "Selecione um jogador na tabela.")
            return

        confirmar = messagebox.askyesno(
            "Confirmacao", "Deseja realmente excluir o jogador selecionado?"
        )
        if not confirmar:
            return

        excluido = excluir_cliente(self.id_cliente)
        if not excluido:
            messagebox.showerror("Erro", "Jogador nao encontrado.")
            return

        self.carregar_clientes()
        self.limpar_campos()
        self.status_var.set("Jogador excluido com sucesso.")
        messagebox.showinfo("Sucesso", "Jogador excluido com sucesso.")

    def selecionar_cliente(self, _evento):
        item_selecionado = self.tabela.selection()
        if not item_selecionado:
            return

        valores = self.tabela.item(item_selecionado[0], "values")
        self.id_cliente = int(valores[0])
        cliente = buscar_cliente_por_id(self.id_cliente)
        if cliente is None:
            return

        self.nome_var.set(cliente[1])
        self.email_var.set(cliente[2])
        self.telefone_var.set(cliente[3])
        self.cpf_var.set(cliente[4])
        self.cidade_var.set(cliente[5])
        self.time_atual_var.set(cliente[6])
        self.pe_dominante_var.set(cliente[7])
        self.numero_camisa_var.set(cliente[8])
        self.posicao_var.set(cliente[9])
        if cliente[9] == "Goleiro":
            self.desempenho_var.set(cliente[11])
        else:
            self.desempenho_var.set(cliente[10])
        self.foto_var.set(cliente[12])
        self.atualizar_preview_foto()

    def editar_por_duplo_clique(self, _evento):
        item_selecionado = self.tabela.selection()
        if not item_selecionado:
            return

        self.selecionar_cliente(None)
        self.nome_entry.focus_set()
        self.status_var.set("Jogador carregado para edicao. Altere os campos e clique em Atualizar.")

    def rolar_tabela_mouse(self, evento):
        self.tabela.yview_scroll(int(-1 * (evento.delta / 120)), "units")
        return "break"

    def limpar_campos(self):
        self.id_cliente = None
        self.nome_var.set("")
        self.email_var.set("")
        self.telefone_var.set("")
        self.cpf_var.set("")
        self.cidade_var.set("")
        self.time_atual_var.set("")
        self.pe_dominante_var.set("")
        self.numero_camisa_var.set("")
        self.posicao_var.set("")
        self.desempenho_var.set("")
        self.foto_var.set("")
        self.atualizar_preview_foto()
        self.tabela.selection_remove(self.tabela.selection())


def main():
    inicializar_planilha()

    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
